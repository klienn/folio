import csv
import io
from contextlib import closing
from pathlib import Path

import pymupdf
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from backend.db import connect
from backend.main import create_app
from backend.worker import process_document, recover_jobs


@pytest.fixture
def workspace(tmp_path):
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        yield client, tmp_path


def make_pdf(*texts):
    pdf = pymupdf.open()
    for text in texts:
        page = pdf.new_page()
        page.insert_text((72, 72), text)
    return pdf.tobytes()


def make_image():
    output = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(output, format="PNG")
    return output.getvalue()


def upload(client, filename="pages.pdf", content=None, **data):
    response = client.post(
        "/api/documents",
        files={"files": (filename, content or make_pdf("First page", "Second page"))},
        data=data,
    )
    assert response.status_code == 201
    return response.json()["documents"][0]


class FakeEngine:
    calls = 0

    def recognize(self, *args):
        self.calls += 1
        return "Text from a scan", 0.91, "CPU"


def test_pdf_review_edits_persist_and_csv_has_one_row_per_page(workspace):
    client, root = workspace
    document = upload(client, filename='a, "quoted".pdf')
    engine = FakeEngine()
    process_document(root, document["id"], engine)
    assert engine.calls == 0
    url = f"/api/documents/{document['id']}"
    detail = client.get(url).json()
    assert detail["status"] == "completed"
    assert [p["text"] for p in detail["pages"]] == ["First page", "Second page"]
    assert client.get(f"{url}/pages/2/preview").headers["content-type"] == "image/png"
    assert client.get(f"{url}/original").content.startswith(b"%PDF")
    text = 'Corrected, "quoted" text\n日本語 — العربية'
    response = client.patch(f"{url}/pages/1", json={"text": text, "revision": 0})
    assert response.status_code == 200
    assert response.json()["raw_text"] == "First page"
    assert response.json()["revision"] == 1
    assert client.patch(f"{url}/pages/1", json={"text": "stale", "revision": 0}).status_code == 409
    # An explicitly empty correction must not revert to the raw extraction.
    assert client.patch(f"{url}/pages/2", json={"text": "", "revision": 0}).status_code == 200
    with TestClient(create_app(root, start_worker=False)) as reopened:
        assert reopened.get(url).json()["pages"][0]["text"] == text
        result = reopened.get("/api/export", params={"ids": document["id"]})
        rows = list(csv.DictReader(io.StringIO(result.content.decode("utf-8-sig"))))
    assert len(rows) == 2
    assert list(rows[0]) == ["document_id", "filename", "page_number", "text", "status"]
    assert rows[0]["filename"] == 'a, "quoted".pdf'
    assert rows[0]["text"] == text
    assert rows[1]["text"] == ""
    assert rows[0]["status"] == "completed"


def test_mixed_pdf_and_force_ocr(workspace):
    client, root = workspace
    pdf = pymupdf.open(stream=make_pdf("Digital text", ""), filetype="pdf")
    pdf[1].insert_image(pymupdf.Rect(40, 40, 300, 300), stream=make_image())
    document = upload(client, content=pdf.tobytes())
    engine = FakeEngine()
    process_document(root, document["id"], engine)
    pages = client.get(f"/api/documents/{document['id']}").json()["pages"]
    assert [p["method"] for p in pages] == ["embedded", "ocr"]
    assert engine.calls == 1
    forced = upload(client, force_ocr="true")
    process_document(root, forced["id"], engine)
    assert engine.calls == 3


def test_image_retry_and_delete(workspace):
    client, root = workspace
    document = upload(client, "scan.png", make_image())
    url = f"/api/documents/{document['id']}"
    assert client.delete(url).status_code == 409

    class BrokenEngine:
        def recognize(self, *args):
            raise RuntimeError("Model download unavailable")

    process_document(root, document["id"], BrokenEngine())
    detail = client.get(url).json()
    assert detail["status"] == "failed"
    assert "download unavailable" in detail["pages"][0]["error"]
    assert client.post(f"{url}/retry").status_code == 200
    process_document(root, document["id"], FakeEngine())
    assert client.get(url).json()["pages"][0]["text"] == "Text from a scan"
    assert client.delete(url).status_code == 204
    assert client.get(url).status_code == 404
    assert not list((root / "uploads").iterdir())
    assert not list((root / "previews").iterdir())


@pytest.mark.parametrize("status", ["completed", "failed"])
def test_delete_removes_all_document_data_and_preserves_other_documents(workspace, status):
    client, root = workspace
    document = upload(client, "delete-me.pdf")
    other = upload(client, "keep-me.pdf")
    for item in (document, other):
        process_document(root, item["id"], FakeEngine())
    url = f"/api/documents/{document['id']}"
    assert (
        client.patch(f"{url}/pages/1", json={"text": "Saved correction", "revision": 0}).status_code
        == 200
    )
    model = root / "models" / "shared-model.onnx"
    model.write_bytes(b"shared model")
    with connect(root) as db:
        original = (
            root
            / db.execute("SELECT path FROM documents WHERE id = ?", (document["id"],)).fetchone()[0]
        )
        db.execute(
            "UPDATE documents SET status = ?, error = 'test error' WHERE id = ?",
            (status, document["id"]),
        )
    assert original.exists()
    assert (root / "previews" / document["id"]).is_dir()

    assert client.delete(url).status_code == 204
    assert not original.exists()
    assert not (root / "previews" / document["id"]).exists()
    with connect(root) as db:
        assert (
            db.execute("SELECT count(*) FROM documents WHERE id = ?", (document["id"],)).fetchone()[
                0
            ]
            == 0
        )
        assert (
            db.execute(
                "SELECT count(*) FROM pages WHERE document_id = ?", (document["id"],)
            ).fetchone()[0]
            == 0
        )
    for suffix in ("", "/original", "/pages/1/preview", "/pages/2/preview"):
        assert client.get(url + suffix).status_code == 404
    assert client.delete(url).status_code == 404
    assert client.get("/api/export", params={"ids": document["id"]}).status_code == 404
    rows = list(csv.DictReader(io.StringIO(client.get("/api/export").content.decode("utf-8-sig"))))
    assert {row["document_id"] for row in rows} == {other["id"]}
    assert client.get(f"/api/documents/{other['id']}/original").status_code == 200
    assert client.get(f"/api/documents/{other['id']}/pages/1/preview").status_code == 200
    assert model.read_bytes() == b"shared model"


@pytest.mark.parametrize("blocked", ["original", "previews"])
def test_delete_cleanup_failure_is_reported_and_can_be_retried(workspace, monkeypatch, blocked):
    client, root = workspace
    document = upload(client)
    process_document(root, document["id"], FakeEngine())
    url = f"/api/documents/{document['id']}"

    def deny(*args, **kwargs):
        raise PermissionError("File is in use")

    with monkeypatch.context() as patch:
        if blocked == "original":
            patch.setattr(Path, "unlink", deny)
        else:
            patch.setattr("backend.main.shutil.rmtree", deny)
        result = client.delete(url)
        assert result.status_code == 500
        assert "Could not fully delete" in result.json()["detail"]
        assert client.get(url).status_code == 200
        with connect(root) as db:
            assert db.execute("SELECT count(*) FROM pages").fetchone()[0] == 2
    assert client.delete(url).status_code == 204
    assert not list((root / "uploads").iterdir())
    assert not list((root / "previews").iterdir())


def test_delete_handles_missing_files_and_rejects_active_jobs(workspace):
    client, root = workspace
    document = upload(client)
    url = f"/api/documents/{document['id']}"
    with connect(root) as db:
        db.execute("UPDATE documents SET status = 'processing'")
        original = root / db.execute("SELECT path FROM documents").fetchone()[0]
    assert client.delete(url).status_code == 409
    assert original.exists()
    with connect(root) as db:
        db.execute("UPDATE documents SET status = 'failed'")
    original.unlink()
    assert client.delete(url).status_code == 204
    assert client.get("/api/documents").json() == []


def test_recovery_keeps_completed_pages_and_edits(workspace):
    client, root = workspace
    document = upload(client)
    with connect(root) as db:
        db.execute("UPDATE documents SET status='processing'")
        db.execute("UPDATE pages SET status='processing' WHERE page_number=2")
        db.execute(
            "UPDATE pages SET status='completed', raw_text='original', edited_text='saved' "
            "WHERE page_number=1"
        )
    recover_jobs(root)
    process_document(root, document["id"], FakeEngine())
    detail = client.get(f"/api/documents/{document['id']}").json()
    assert detail["status"] == "completed"
    assert detail["pages"][0]["text"] == "saved"
    assert detail["pages"][1]["text"] == "Second page"


def test_upload_validation_and_partial_batch_success(workspace):
    client, root = workspace
    response = client.post(
        "/api/documents",
        files=[
            ("files", ("broken.pdf", b"not a PDF")),
            ("files", ("empty.png", b"")),
            ("files", ("notes.txt", b"unsupported")),
            ("files", ("../valid.png", make_image())),
        ],
    )
    data = response.json()
    assert len(data["documents"]) == 1
    assert data["documents"][0]["filename"] == "valid.png"
    assert len(data["errors"]) == 3
    assert len(list((root / "uploads").iterdir())) == 1
    assert (
        client.post(
            "/api/documents", files={"files": ("a.png", make_image())}, data={"language": "unknown"}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/documents",
            files={"files": ("a.png", make_image())},
            data={"language": "arabic", "handwriting": "true"},
        ).status_code
        == 422
    )


def test_encrypted_pdf_and_size_limits(workspace, monkeypatch):
    client, _ = workspace
    pdf = pymupdf.open(stream=make_pdf("Locked"), filetype="pdf")
    locked = pdf.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    result = client.post("/api/documents", files={"files": ("locked.pdf", locked)}).json()
    assert "Password-protected" in result["errors"][0]["error"]
    monkeypatch.setattr("backend.main.MAX_UPLOAD_BYTES", 10)
    result = client.post("/api/documents", files={"files": ("large.png", make_image())}).json()
    assert "50 MB" in result["errors"][0]["error"]


def test_exports_selected_documents_and_neutralizes_formulas(workspace):
    client, root = workspace
    one = upload(client, "=formula.pdf", make_pdf("=1+1"))
    upload(client, "other.pdf")
    process_document(root, one["id"], FakeEngine())
    export = client.get("/api/export", params={"ids": one["id"]})
    rows = list(csv.DictReader(io.StringIO(export.content.decode("utf-8-sig"))))
    assert len(rows) == 1
    assert rows[0]["text"] == "'=1+1"
    assert rows[0]["filename"] == "'=formula.pdf"
    assert client.get("/api/export", params={"ids": "nonexistent"}).status_code == 404


def test_excel_export_preserves_edits_and_formats_pages(workspace):
    client, root = workspace
    document = upload(client, filename='a, "quoted".pdf')
    upload(client, filename="unselected.pdf")
    process_document(root, document["id"], FakeEngine())
    text = '=Not a formula\nCorrected, "quoted" text\n日本語 — العربية\n' + "Details " * 45
    for number, value in enumerate((text, ""), 1):
        response = client.patch(
            f"/api/documents/{document['id']}/pages/{number}",
            json={"text": value, "revision": 0},
        )
        assert response.status_code == 200
    response = client.get("/api/export", params={"ids": document["id"], "format": "xlsx"})
    assert response.status_code == 200
    assert response.headers["content-type"].endswith("spreadsheetml.sheet")
    assert 'filename="a, -quoted--ocr_' in response.headers["content-disposition"]
    assert '.xlsx";' in response.headers["content-disposition"]
    with closing(load_workbook(io.BytesIO(response.content))) as workbook:
        sheet = workbook["Pages"]
        assert sheet.max_row == 3
        assert sheet.max_column == 5
        assert list(next(sheet.values)) == [
            "document_id",
            "filename",
            "page_number",
            "text",
            "status",
        ]
        assert sheet["A2"].value == document["id"]
        assert sheet["B2"].value == 'a, "quoted".pdf'
        assert sheet["C2"].value == 1
        assert sheet["C2"].data_type == "n"
        assert sheet["D2"].value == text
        assert sheet["D2"].data_type == "s"
        assert sheet["D3"].value in (None, "")
        assert sheet["E2"].value == "completed"
        assert sheet["D2"].alignment.wrap_text
        assert sheet["D2"].alignment.vertical == "top"
        assert sheet["A1"].font.bold
        assert sheet["D2"].font.name == "Arial"
        assert sheet.column_dimensions["D"].width >= 80
        assert sheet.row_dimensions[2].height > sheet.row_dimensions[3].height
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:E3"
        assert all(cell.data_type != "f" for row in sheet for cell in row)


def test_excel_all_pages_and_empty_workspace(workspace):
    client, root = workspace
    empty = client.get("/api/export", params={"format": "xlsx"})
    assert 'filename="folio-library-0-documents_' in empty.headers["content-disposition"]
    with closing(load_workbook(io.BytesIO(empty.content))) as workbook:
        assert workbook.active.max_row == 1
    one = upload(client, filename="=formula.pdf", content=make_pdf("=1+1"))
    upload(client, filename="queued.pdf", content=make_pdf("Pending"))
    process_document(root, one["id"], FakeEngine())
    with connect(root) as db:
        db.execute(
            "UPDATE pages SET status = 'failed', error = 'Private diagnostic' "
            "WHERE document_id = ?",
            (one["id"],),
        )
    response = client.get("/api/export", params={"format": "xlsx"})
    assert 'filename="folio-library-2-documents_' in response.headers["content-disposition"]
    with closing(load_workbook(io.BytesIO(response.content))) as workbook:
        rows = list(workbook.active.values)
        assert len(rows) == 3
        assert rows[1][1:] == ("=formula.pdf", 1, "=1+1", "failed")
        assert rows[2][4] == "pending"
    assert client.get("/api/export", params={"ids": "missing", "format": "xlsx"}).status_code == 404
    assert client.get("/api/export", params={"format": "unsupported"}).status_code == 422


def test_excel_rejects_oversized_cells_without_truncating_csv(workspace):
    client, root = workspace
    document = upload(client, content=make_pdf("Page"))
    process_document(root, document["id"], FakeEngine())
    text = "a" * 32768
    client.patch(f"/api/documents/{document['id']}/pages/1", json={"text": text, "revision": 0})
    response = client.get("/api/export", params={"format": "xlsx"})
    assert response.status_code == 422
    assert "32,767" in response.json()["detail"]
    csv_response = client.get("/api/export")
    rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert rows[0]["text"] == text


def test_cross_origin_changes_rejected(workspace):
    client, _ = workspace
    response = client.post("/api/documents", headers={"origin": "https://other.example"})
    assert response.status_code == 403


def test_unfinished_page_cannot_be_edited(workspace):
    client, _ = workspace
    document = upload(client)
    assert (
        client.patch(
            f"/api/documents/{document['id']}/pages/1", json={"text": "edit", "revision": 0}
        ).status_code
        == 409
    )
    assert client.get(f"/api/documents/{document['id']}/pages/3/preview").status_code == 404
